"""Source-parameter response twins and a paired-report design diagnostic.

This is a mathematical instantiation at released fitted parameters, not new
human data, refitting, a power analysis, or evidence for criterion variability.
Gaussian noise decomposition and paired-noise methods have prior literature.
The source experiment observed separate task blocks, so all joint predictions
below are explicitly prospective and depend on a shared sensory draw.
"""
from pathlib import Path
import argparse
import csv
import hashlib
import json
import math

import numpy as np
from openpyxl import load_workbook
from scipy.integrate import quad
from scipy.special import ndtr, xlogy

HERE = Path(__file__).resolve().parent
DEFAULT_SOURCE = HERE.parent / "empirical" / "inputs"
SOA = np.array([-400., -200., -100., 0., 100., 200., 400.])


def threshold(prior, variance, stimulus_variance):
    k2 = 2 * variance * (variance + stimulus_variance) / stimulus_variance * (
        np.log(prior) - np.log1p(-prior)
        + .5 * np.log1p(stimulus_variance / variance)
    )
    return float(np.sqrt(max(0., k2)))


def interval_probability(soa, k, variance):
    sd = np.sqrt(variance)
    return ndtr((k - soa) / sd) - ndtr((-k - soa) / sd)


def centered_rectangle(r1, r2, rho):
    """Deterministic one-dimensional integration; no randomized CDF call."""
    if r1 == 0 or r2 == 0:
        return 0.
    if rho >= 1:
        return float(2 * ndtr(min(r1, r2)) - 1)
    if rho <= 0:
        return float((2 * ndtr(r1) - 1) * (2 * ndtr(r2) - 1))
    den = np.sqrt(1 - rho * rho)
    def integrand(z):
        return np.exp(-.5*z*z)/np.sqrt(2*np.pi) * (
            ndtr((r2-rho*z)/den) - ndtr((-r2-rho*z)/den))
    value, error = quad(integrand, -r1, r1, epsabs=2e-12, epsrel=2e-12, limit=200)
    if error > 1e-9:
        raise RuntimeError(f"Rectangle integral error {error}")
    return float(value)


def observed_joint(q1, q2, q12, lapse):
    # Independent lapse decisions and independent fair guesses in the two tasks.
    p1 = lapse/2 + (1-lapse)*q1
    p2 = lapse/2 + (1-lapse)*q2
    p11 = (1-lapse)**2*q12 + lapse*(1-lapse)*(q1+q2)/2 + lapse**2/4
    result = np.array([p11, p1-p11, p2-p11, 1-p1-p2+p11])
    if np.min(result) < -1e-12 or abs(np.sum(result)-1) > 1e-10:
        raise ValueError(result)
    # Only erase roundoff-size negatives, not fit probabilities.
    result=np.maximum(result,0.)
    return result/result.sum()


def save_csv(path, rows):
    with path.open('w', newline='') as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]))
        writer.writeheader(); writer.writerows(rows)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--source', type=Path, default=DEFAULT_SOURCE)
    args = parser.parse_args()
    model_path = args.source / 'Experiment_3_Computational_modelling.xlsx'
    counts_path = args.source / 'Experiment_3.xlsx'
    models = load_workbook(model_path, data_only=True).active
    counts_sheet = load_workbook(counts_path, data_only=True).active
    rows, joint_rows, people = [], [], []
    for i in range(30):
        pars = np.array([models.cell(i+36, j).value for j in range(2,9)], float)
        priors = pars[:2]; variances = np.exp(2*pars[2:5])
        stimulus_var = float(np.exp(2*pars[5])); lapse = float(pars[6])
        # Fixed construction rule; no data-dependent selection for separation.
        constant_sensory_variance = .5 * float(np.min(variances))
        data = np.array([counts_sheet.cell(i+4, j+2).value for j in range(42)], float).reshape(2,3,7)
        nll_a, nll_b, max_gap = 0., 0., 0.
        for f, label in enumerate(['8Hz', 'sham', '13Hz']):
            v = variances[f]
            jitter_var = v - constant_sensory_variance
            k = np.array([threshold(p, v, stimulus_var) for p in priors])
            q_a = np.stack([interval_probability(SOA, kt, v) for kt in k])
            q_b = np.stack([interval_probability(SOA, kt, constant_sensory_variance+jitter_var) for kt in k])
            p_a = lapse/2 + (1-lapse)*q_a
            p_b = lapse/2 + (1-lapse)*q_b
            max_gap = max(max_gap, float(np.max(np.abs(p_a-p_b))))
            nll_a -= float(np.sum(xlogy(data[:,f], p_a)+xlogy(10-data[:,f], 1-p_a)))
            nll_b -= float(np.sum(xlogy(data[:,f], p_b)+xlogy(10-data[:,f], 1-p_b)))
            for t, task in enumerate(['ownership','simultaneity']):
                for j,s in enumerate(SOA):
                    rows.append({'participant':i+1,'task':task,'condition':label,'soa_ms':s,
                        'released_effective_variance':v,'sensory_variance_model_a':v,
                        'sensory_variance_model_b':constant_sensory_variance,
                        'criterion_center_variance_model_b':jitter_var,
                        'threshold_ms':k[t],'lapse':lapse,'p_a':p_a[t,j],'p_b':p_b[t,j],
                        'yes':int(data[t,f,j]),'n':10})
            r1,r2 = k/np.sqrt(v)
            rho_b = constant_sensory_variance/v
            qa11 = centered_rectangle(r1,r2,1.)
            qb11 = centered_rectangle(r1,r2,rho_b)
            obs_a = observed_joint(q_a[0,3],q_a[1,3],qa11,lapse)
            obs_b = observed_joint(q_b[0,3],q_b[1,3],qb11,lapse)
            # A direct affinity sum rounds to one for nearly saturated profiles.
            # Compute 1-affinity from squared root differences instead.
            hellinger_squared=float(.5*np.sum((np.sqrt(obs_a)-np.sqrt(obs_b))**2))
            affinity = 1-hellinger_squared
            if hellinger_squared == 0:
                n_bound = None
            else:
                # Equal-prior simple-model Bayes error <= affinity**m/2.
                # Not a confidence interval or frequentist power guarantee.
                n_bound = math.ceil(np.log(.1)/np.log1p(-hellinger_squared))
            kl_ab = float(np.sum(xlogy(obs_a, np.divide(obs_a,obs_b,out=np.ones(4),where=obs_b>0))))
            if np.any((obs_a>0)&(obs_b==0)):
                kl_ab = math.inf
            joint_rows.append({'participant':i+1,'condition':label,'r_ownership':r1,'r_simultaneity':r2,
                'lapse':lapse,'rho_model_a':1.,'rho_model_b':rho_b,
                'yes_ownership':float(lapse/2+(1-lapse)*q_a[0,3]),
                'yes_simultaneity':float(lapse/2+(1-lapse)*q_a[1,3]),
                'paired_yes_yes_a':obs_a[0],'paired_yes_yes_b':obs_b[0],
                'paired_yes_yes_difference':obs_a[0]-obs_b[0],
                'affinity_rounded':affinity,'hellinger_squared':hellinger_squared,'kl_a_to_b_nats':kl_ab,
                'known_simple_pair_error_05_sufficient_m':n_bound,
                **{'p_a_'+c:obs_a[j] for j,c in enumerate(['11','10','01','00'])},
                **{'p_b_'+c:obs_b[j] for j,c in enumerate(['11','10','01','00'])}})
        saved_nll = float(models.cell(i+36,9).value)
        people.append({'participant':i+1,'nll_a':nll_a,'nll_b':nll_b,'released_nll':saved_nll,
            'nll_difference':nll_b-nll_a,'max_probability_difference':max_gap,
            'constant_sensory_variance_b':constant_sensory_variance,
            'effective_sigma_8':np.sqrt(variances[0]),'effective_sigma_sham':np.sqrt(variances[1]),
            'effective_sigma_13':np.sqrt(variances[2])})
    save_csv(HERE/'response_twin_probabilities.csv',rows)
    save_csv(HERE/'paired_readout_predictions.csv',joint_rows)
    save_csv(HERE/'response_twin_people.csv',people)
    gaps=np.array([x['paired_yes_yes_difference'] for x in joint_rows])
    bounds=sorted([x['known_simple_pair_error_05_sufficient_m'] for x in joint_rows if x['known_simple_pair_error_05_sufficient_m'] is not None])
    summary={'n_participants':30,'response_cells':len(rows),'prospective_joint_profiles':len(joint_rows),
        'construction':'a_i = 0.5 min_f sigma_if^2; criterion jitter variance = sigma_if^2-a_i; effective thresholds unchanged',
        'max_response_probability_discrepancy':max(x['max_probability_difference'] for x in people),
        'max_nll_discrepancy':max(abs(x['nll_difference']) for x in people),
        'max_saved_nll_discrepancy':max(abs(x['nll_a']-x['released_nll']) for x in people),
        'sum_nll_a':sum(x['nll_a'] for x in people),'sum_nll_b':sum(x['nll_b'] for x in people),
        'paired_yes_yes_gap':{'minimum':float(gaps.min()),'median':float(np.median(gaps)),'maximum':float(gaps.max())},
        'known_simple_pair_error_05_bound':{'finite_profiles':len(bounds),'minimum':min(bounds),'median':float(np.median(bounds)),'maximum':max(bounds)},
        'affinity_numerical_note':'Use h2=.5 sum(sqrt(P)-sqrt(Q))^2 and log1p(-h2). Direct affinity summation rounded three weakly separated profiles to one; that did not mean mathematical nonidentification.',
        'calibration_required':['marginal effective widths and thresholds known','lapse known','one shared sensory sample on a paired trial','independent Gaussian criterion jitters across tasks and from sensory noise','independent lapse and fair guesses','fixed zero-SOA centers'],
        'not_observed_data':'Joint probabilities and count bounds are model-derived prospective diagnostics, not actual joint judgments, trial-level data or a validated experimental power calculation.',
        'known_model_bound':'For two fully specified categorical joint distributions with equal prior odds, optimal average error after m independent paired trials is <= affinity^m/2. Estimated parameters, composite alternatives and calibration error invalidate interpreting this number as actual required sample size.',
        'not_a_new_general_noise_decomposition':'Credit Gaussian criterion-noise and multipass literature; this script instantiates the equivalence at the released intervention parameters.',
        'sources':[{'file':p.name,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in [model_path,counts_path]]}
    assert summary['max_response_probability_discrepancy'] < 1e-12
    assert summary['max_nll_discrepancy'] < 1e-9
    assert summary['max_saved_nll_discrepancy'] < 1e-8
    (HERE/'response_twin_summary.json').write_text(json.dumps(summary,indent=2)+'\n')
    print(json.dumps(summary,indent=2))


if __name__=='__main__':
    main()
