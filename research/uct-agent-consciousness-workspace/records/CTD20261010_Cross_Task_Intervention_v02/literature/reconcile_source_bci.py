"""Read-only reconciliation of released BCI parameters, NLLs and Fig. S2.

No optimization and no relabeling. The model-prediction expression is transcribed
from the authors' modelprediction_log_BCI.m and the two NLL helpers. All quantities
are evaluated on the 30 participant rows in their original order.
"""
from pathlib import Path
import csv
import hashlib
import json

import numpy as np
from openpyxl import load_workbook
from scipy.special import ndtr, xlogy

HERE = Path(__file__).resolve().parent
SOURCE = HERE.parent / "empirical" / "inputs"


def source_probability(s, prior, logsigma, logsigma_s, lapse):
    sigma2 = np.exp(2 * logsigma)
    sigma_s2 = np.exp(2 * logsigma_s)
    threshold2 = (2 * sigma2 * (sigma2 + sigma_s2) / sigma_s2) * (
        np.log(prior) - np.log1p(-prior)
        + .5 * np.log((sigma2 + sigma_s2) / sigma2)
    )
    if threshold2 < 0:
        pc1 = np.zeros_like(s, dtype=float)
    else:
        threshold = np.sqrt(threshold2)
        sigma = np.exp(logsigma)
        pc1 = ndtr((threshold - s) / sigma) - ndtr((-threshold - s) / sigma)
    return .5 * lapse + (1 - lapse) * pc1


def nll_at_saved(parameters, counts, model):
    s = np.array([-400, -200, -100, 0, 100, 200, 400], dtype=float)
    probability = np.empty((6, 7))
    for row in range(6):
        if model == "sigma":
            prior = parameters[row // 3]
            logsigma = parameters[2 + row % 3]
            logsigma_s, lapse = parameters[5:7]
        elif model == "psame":
            prior = parameters[row]
            logsigma, logsigma_s, lapse = parameters[6:9]
        else:
            raise ValueError(model)
        probability[row] = source_probability(s, prior, logsigma, logsigma_s, lapse)
    return float(-np.sum(xlogy(counts, probability) + xlogy(10 - counts, 1 - probability)))


def main():
    model_file = SOURCE / "Experiment_3_Computational_modelling.xlsx"
    count_file = SOURCE / "Experiment_3.xlsx"
    publisher_file = SOURCE / "NatComm2026_SourceData.xlsx"
    model_sheet = load_workbook(model_file, data_only=True).active
    count_sheet = load_workbook(count_file, data_only=True).active
    figure_sheet = load_workbook(publisher_file, data_only=True)["Fig S2"]
    rows = []
    for i in range(30):
        counts = np.array([[count_sheet.cell(i + 4, j + 2).value for j in range(42)]], dtype=float).reshape(6, 7)
        psame_pars = np.array([model_sheet.cell(i + 2, j).value for j in range(2, 11)], dtype=float)
        sigma_pars = np.array([model_sheet.cell(i + 36, j).value for j in range(2, 9)], dtype=float)
        saved_p = float(model_sheet.cell(i + 2, 11).value)
        saved_s = float(model_sheet.cell(i + 36, 9).value)
        eval_p = nll_at_saved(psame_pars, counts, "psame")
        eval_s = nll_at_saved(sigma_pars, counts, "sigma")
        delta_nll = saved_s - saved_p
        rows.append({
            "participant": i + 1,
            "saved_nll_sigma": saved_s, "saved_nll_psame": saved_p,
            "recomputed_nll_sigma": eval_s, "recomputed_nll_psame": eval_p,
            "sigma_nll_error": eval_s - saved_s, "psame_nll_error": eval_p - saved_p,
            "correct_delta_aic": -4 + 2 * delta_nll,
            "public_script_delta_aic": -4 - 2 * delta_nll,
            "publisher_delta_aic": float(figure_sheet.cell(i + 2, 1).value),
            "correct_delta_bic": -2 * np.log(420) + 2 * delta_nll,
            "public_script_delta_bic": -2 * np.log(420) - 2 * delta_nll,
            "publisher_delta_bic": float(figure_sheet.cell(i + 2, 2).value),
        })
    out = HERE / "source_bci_reconciliation.csv"
    with out.open("w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0]))
        writer.writeheader(); writer.writerows(rows)
    summary = {
        "inputs": [{"path": str(p), "sha256": hashlib.sha256(p.read_bytes()).hexdigest()}
                   for p in [model_file, count_file, publisher_file]],
        "n": 30, "delta_convention": "BCI_sigma minus BCI_psame; negative favors sigma",
        "free_parameter_difference": -2, "trials_per_person": 420,
        "no_optimization_performed": True, "no_participant_permutation": True,
        "max_abs_saved_parameter_nll_error": {
            model: max(abs(r[model + "_nll_error"]) for r in rows) for model in ["sigma", "psame"]},
        "max_abs_publisher_minus_script": {
            ic: max(abs(r["publisher_delta_" + ic] - r["public_script_delta_" + ic]) for r in rows)
            for ic in ["aic", "bic"]},
        "max_abs_publisher_minus_correct": {
            ic: max(abs(r["publisher_delta_" + ic] - r["correct_delta_" + ic]) for r in rows)
            for ic in ["aic", "bic"]},
        "sums": {key: sum(r[key] for r in rows) for key in [
            "saved_nll_sigma", "saved_nll_psame", "correct_delta_aic", "public_script_delta_aic",
            "publisher_delta_aic", "correct_delta_bic", "public_script_delta_bic", "publisher_delta_bic"]},
        "raw_nll_favor_sigma_n": int(sum(r["saved_nll_sigma"] < r["saved_nll_psame"] for r in rows)),
        "correct_aic_favor_sigma_n": int(sum(r["correct_delta_aic"] < 0 for r in rows)),
        "correct_bic_favor_sigma_n": int(sum(r["correct_delta_bic"] < 0 for r in rows)),
        "interpretation": "The information-criterion sign discrepancy changes reported values, not the aggregate preferred model. These are saved-fit evaluations, not guaranteed global optima or new model fits.",
    }
    (HERE / "source_bci_reconciliation.json").write_text(json.dumps(summary, indent=2) + "\n")
    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
