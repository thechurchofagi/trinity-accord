"""Independent source BCI likelihood with optional stable task-specific centers.

Model families: sigma / prior, and the same two models with two task centers.
Center parameters are in 100-ms units internally. Effective parameter counts
are 6, 8, 8, 10; the stimulus prior scale is fixed, not a free parameter.
No CV split may use published full-data parameters as an initialization.
"""
from pathlib import Path
import argparse
import json

import numpy as np
from scipy.optimize import minimize
from scipy.special import ndtr, logit, expit, xlogy
from scipy.stats import qmc

SOA = np.array([-400.,-200.,-100.,0.,100.,200.,400.])
STIMULUS_VARIANCE = float(np.var([-400.,-200.,-100.,100.,200.,400.], ddof=1))
MODELS = ('sigma','prior','sigma_shift','prior_shift')


def structure(model):
    if model not in MODELS:
        raise ValueError(model)
    is_sigma = model.startswith('sigma')
    nu, ns = (2,3) if is_sigma else (6,1)
    ui = np.repeat(np.arange(2),3) if is_sigma else np.arange(6)
    si = np.tile(np.arange(3),2) if is_sigma else np.zeros(6,dtype=int)
    shift = model.endswith('_shift')
    return nu,ns,ui,si,shift


def bounds(model):
    nu,ns,_,_,shift=structure(model)
    return [(-16.,16.)]*nu+[(-10.,10.)]*ns+[(1e-8,1-1e-8)]+([(-8.,8.)]*2 if shift else [])


def probabilities_and_derivatives(theta, model):
    theta=np.asarray(theta,float)
    nu,ns,ui,si,shift=structure(model)
    u=theta[ui]
    logsd=theta[nu+si]
    variance=np.exp(2*logsd)
    sd=np.exp(logsd)
    lapse=theta[nu+ns]
    center=np.repeat(theta[-2:]*100.,3) if shift else np.zeros(6)
    relative=SOA[None,:]-center[:,None]
    absolute=np.abs(relative)
    a=2*variance*(variance+STIMULUS_VARIANCE)/STIMULUS_VARIANCE
    b=u+.5*np.log1p(STIMULUS_VARIANCE/variance)
    k2=a*b
    positive=k2>0
    k=np.sqrt(np.maximum(k2,0.))
    safe_k=np.where(positive,k,1.)
    z1=(k[:,None]-absolute)/sd[:,None]
    z0=(-k[:,None]-absolute)/sd[:,None]
    q=ndtr(z1)-ndtr(z0)
    phi1=np.exp(-.5*z1*z1)/np.sqrt(2*np.pi)
    phi0=np.exp(-.5*z0*z0)/np.sqrt(2*np.pi)
    dq_dk=(phi1+phi0)/sd[:,None]
    dk_du=np.where(positive,a/(2*safe_k),0.)
    dk_dlogsigma=np.where(positive,((4*variance+8*variance*variance/STIMULUS_VARIANCE)*b-2*variance)/(2*safe_k),0.)
    dq_du=dq_dk*dk_du[:,None]
    dq_dlogsigma=(phi0*z0-phi1*z1)+dq_dk*dk_dlogsigma[:,None]
    dq_dlogsigma[~positive]=0.
    dq_dcenter=np.sign(relative)*(phi1-phi0)/sd[:,None]
    probability=lapse/2+(1-lapse)*q
    return probability,q,dq_du,dq_dlogsigma,dq_dcenter


def nll_and_gradient(theta, counts, trials, model):
    counts=np.asarray(counts,float).reshape(6,7)
    trials=np.broadcast_to(np.asarray(trials,float),counts.shape)
    p,q,du,ds,dm=probabilities_and_derivatives(theta,model)
    # Bounds on the explicit lapse keep p strictly interior, without clipping p.
    if np.any(p<=0) or np.any(p>=1):
        raise FloatingPointError('Probability outside strict interior')
    nll=-float(np.sum(xlogy(counts,p)+xlogy(trials-counts,1-p)))
    w=(trials*p-counts)/(p*(1-p))
    nu,ns,ui,si,shift=structure(model)
    lapse=theta[nu+ns]
    grad=np.zeros(len(theta))
    np.add.at(grad,ui,np.sum(w*(1-lapse)*du,axis=1))
    np.add.at(grad,nu+si,np.sum(w*(1-lapse)*ds,axis=1))
    grad[nu+ns]=np.sum(w*(.5-q))
    if shift:
        center_rows=np.sum(w*(1-lapse)*dm*100.,axis=1)
        grad[-2:]=center_rows.reshape(2,3).sum(axis=1)
    return nll,grad


def generic_starts(model, n_starts, seed):
    nu,ns,_,_,shift=structure(model)
    dim=nu+ns+1+(2 if shift else 0)
    exponent=int(np.ceil(np.log2(max(1,n_starts))))
    unit=qmc.Sobol(d=dim,scramble=True,seed=seed).random_base2(exponent)[:n_starts]
    result=np.empty_like(unit)
    result[:,:nu]=logit(.25+.7*unit[:,:nu])
    result[:,nu:nu+ns]=np.log(35.)+unit[:,nu:nu+ns]*np.log(800./35.)
    result[:,nu+ns]=.005+.245*unit[:,nu+ns]
    if shift:
        result[:,-2:]=-.4+.8*unit[:,-2:]
    canonical=np.r_[np.full(nu,logit(.65)),np.full(ns,np.log(180.)),.04,([0.,0.] if shift else [])]
    return np.vstack([canonical,result])


def source_initial(parameters, model):
    p=np.asarray(parameters,float)
    if model.startswith('sigma'):
        theta=np.r_[logit(p[:2]),p[2:5],p[6]]
    else:
        theta=np.r_[logit(p[:6]),p[6],p[8]]
    if model.endswith('_shift'):
        theta=np.r_[theta,0.,0.]
    bs=np.array(bounds(model));return np.clip(theta,bs[:,0],bs[:,1])


def projected_gradient(theta, gradient, bs):
    g=np.array(gradient,copy=True);x=np.asarray(theta);b=np.array(bs)
    g[(x<=b[:,0]+1e-7)&(g>0)]=0
    g[(x>=b[:,1]-1e-7)&(g<0)]=0
    return float(np.max(np.abs(g)))


def fit_one(counts, trials, model, seed=20261010, n_starts=16, additional_starts=None):
    starts=list(generic_starts(model,n_starts,seed))
    if additional_starts is not None:
        starts.extend(np.asarray(additional_starts,float))
    bs=bounds(model);results=[]
    for initial in starts:
        res=minimize(nll_and_gradient,initial,args=(counts,trials,model),jac=True,
            method='L-BFGS-B',bounds=bs,
            options={'maxiter':1500,'maxls':50,'ftol':1e-12,'gtol':1e-7})
        f,g=nll_and_gradient(res.x,counts,trials,model)
        results.append((f,res.x.copy(),bool(res.success),int(res.nit),str(res.message),projected_gradient(res.x,g,bs)))
    best=min(results,key=lambda r:r[0])
    if not best[2] or best[5]>1e-3:
        # A targeted numerical refinement, applied by a fixed rule, not p value.
        res=minimize(nll_and_gradient,best[1],args=(counts,trials,model),jac=True,
            method='L-BFGS-B',bounds=bs,
            options={'maxiter':4000,'maxls':100,'ftol':1e-14,'gtol':1e-8})
        f,g=nll_and_gradient(res.x,counts,trials,model)
        results.append((f,res.x.copy(),bool(res.success),int(res.nit),str(res.message),projected_gradient(res.x,g,bs)))
        best=min(results,key=lambda r:r[0])
    return {'nll':best[0],'theta':best[1],'success':best[2],
        'projected_gradient_max':best[5],'n_starts':len(results),
        'successful_starts':sum(x[2] for x in results),
        'start_nlls':np.array([x[0] for x in results]),
        'all_start_success':np.array([x[2] for x in results]),
        'best_message':best[4],
        'bounds_hit':[i for i,(v,b) in enumerate(zip(best[1],bs)) if min(abs(v-b[0]),abs(v-b[1]))<1e-6]}


def fit_summary(fit,model):
    nu,ns,_,_,shift=structure(model)
    x=fit['theta']
    return {'priors':expit(x[:nu]).tolist(),'sigma_ms':np.exp(x[nu:nu+ns]).tolist(),
        'lapse':float(x[nu+ns]),'task_centers_ms':(100*x[-2:]).tolist() if shift else [0.,0.]}


def self_check():
    from scipy.optimize._numdiff import approx_derivative
    from openpyxl import load_workbook
    source=Path(__file__).resolve().parents[2]/'empirical'
    workbook=load_workbook(source/'Experiment_3.xlsx',data_only=True).active
    models=load_workbook(source/'Experiment_3_Computational_modelling.xlsx',data_only=True).active
    y=np.array([workbook.cell(4,j+2).value for j in range(42)],float).reshape(6,7)
    checks=[]
    for mi,model in enumerate(MODELS):
        initial=generic_starts(model,4,631+mi)
        errors=[]
        for x in initial:
            analytical=nll_and_gradient(x,y,10,model)[1]
            numeric=approx_derivative(lambda z:nll_and_gradient(z,y,10,model)[0],x,method='3-point').ravel()
            errors.append(float(np.max(np.abs(analytical-numeric)/(1+np.abs(numeric)))))
        assert max(errors)<2e-4,(model,errors)
        checks.append({'model':model,'max_scaled_gradient_error':max(errors)})
    for base,row,cols,nll_col in [('sigma',36,range(2,9),9),('prior',2,range(2,11),11)]:
        parameters=np.array([models.cell(row,j).value for j in cols],float)
        x=source_initial(parameters,base)
        value=nll_and_gradient(x,y,10,base)[0]
        saved=float(models.cell(row,nll_col).value)
        assert abs(value-saved)<1e-5,(base,value,saved)
        shifted=np.r_[x,0.,0.]
        assert abs(nll_and_gradient(shifted,y,10,base+'_shift')[0]-value)<1e-9
        checks.append({'source_model':base,'saved_nll':saved,'independent_nll':value,'difference':value-saved})
    report={'checks':checks,'probability_clipping':False,'lapse_floor_is_explicit_parameter_bound':1e-8,
        'prospective_cv_policy':'generic starts only; never source saved full-data estimates',
        'center_bounds_ms':[-800,800],'sigma_log_bounds':[-10,10],'prior_logit_bounds':[-16,16],
        'source_fixed_stimulus_variance':STIMULUS_VARIANCE}
    (Path(__file__).resolve().parent/'response_fitter_self_check.json').write_text(json.dumps(report,indent=2)+'\n')
    print(json.dumps(report,indent=2))


if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('--self-check',action='store_true')
    args=parser.parse_args()
    if args.self_check:self_check()
