#!/usr/bin/env python3
"""Build and seal the reviewed conceptual paper, not a scientific experiment.

Run build with the primary runtime after the artifact-operation markers.
Render both DOCX files and inspect every page before finalize --visual-review-confirmed.
This program never reserves, uploads or publishes a DOI.
"""
from __future__ import annotations

import argparse
import hashlib
import html
from html.parser import HTMLParser
import json
from pathlib import Path
import re
import unicodedata

from docx import Document
from docx.shared import Pt
import fitz

import document_helpers as documents
from publication_common import ROOT, REPORT, VERSION, TITLE, DATE, STEM, ALLOWED_FILES, extract_abstract, validate_identity
from check_site_contract import HeadMetadata

BASE = "https://www.trinityaccord.org/research/learning-from-an-ai-claimant/"
LANGS = {"en": ("EN.md", ""), "zh": ("ZH_Guide.md", "-zh-guide")}
CSS = """*{box-sizing:border-box}body{font-family:Georgia,'Noto Serif CJK SC',serif;max-width:800px;margin:40px auto;padding:0 22px;line-height:1.7;color:#191919;background:#fff}h1,h2,h3{color:#000;line-height:1.3}h1{font-size:2rem}h2{margin-top:1.7em;font-size:1.4rem}h3{margin-top:1.4em;font-size:1.12rem}header{text-align:center}.subtitle{font-size:1.2rem}.frontmatter{font-size:.95rem}.reference{padding-left:2em;text-indent:-2em;font-size:.94rem}a{color:#194e77;overflow-wrap:anywhere}nav{font:.9rem/1.6 system-ui,sans-serif;margin-bottom:2em}p{overflow-wrap:break-word}@media(max-width:600px){body{margin:22px auto;padding:0 18px;font-size:17px}h1{font-size:1.65rem}}@media print{nav{display:none}body{max-width:none;margin:0}a{color:inherit}}"""


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def stem(lang):
    return STEM + LANGS[lang][1] + "-v" + VERSION


def inline(content):
    parts, cursor = [], 0
    for match in documents.INLINE_RE.finditer(content):
        parts.append(html.escape(content[cursor:match.start()]))
        label, target, raw, bold, italic, code = match.groups()
        if target or raw:
            suffix = ""
            if raw:
                target, suffix = documents.split_url_punctuation(raw)
                label = target
            parts.append(f'<a href="{html.escape(target, quote=True)}">{html.escape(label)}</a>' + html.escape(suffix))
        else:
            tag = "strong" if bold else "em" if italic else "code"
            parts.append(f"<{tag}>{html.escape(bold or italic or code)}</{tag}>")
        cursor = match.end()
    parts.append(html.escape(content[cursor:]))
    return "".join(parts)


class TextOnly(HTMLParser):
    def __init__(self, article_only=False):
        super().__init__(convert_charrefs=True)
        self.parts = []
        self.active = not article_only
        self.article_only = article_only

    def handle_starttag(self, tag, attrs):
        if tag == "article":
            self.active = True

    def handle_endtag(self, tag):
        if self.article_only and tag == "article":
            self.active = False

    def handle_data(self, data):
        if self.active:
            self.parts.append(data)


def plain(value):
    parser = TextOnly()
    parser.feed(inline(value))
    return "".join(parser.parts)


def normalized(value):
    # Ignore typography only; retain every letter and digit in order.
    return "".join(c for c in unicodedata.normalize("NFKC", value) if c.isalnum())


def publication_source(path, deposit):
    text = path.read_text(encoding="utf-8")
    if any(token in text for token in ("TODO", "TBD", "PLACEHOLDER", "")):
        raise ValueError("Unfinished source marker")
    front, body = re.split(r"(?m)(?=^### (?:Abstract|摘要)$)", text, maxsplit=1)
    if re.search(r"(?m)^DOI:", front):
        raise ValueError("Source already has a DOI; review explicitly before rebuilding")
    return front.rstrip() + "\n\nDOI: https://doi.org/" + deposit["doi"] + "\n\n" + body


def render_html(path, deposit, lang):
    blocks = documents.read_blocks(path)
    canonical = BASE if lang == "en" else BASE + "zh-guide.html"
    title = TITLE if lang == "en" else blocks[0][1] + " 中文论证说明"
    metadata = {"citation_title": TITLE, "citation_author": "Hongju Liu", "citation_publication_date": DATE.replace("-", "/"),
                "citation_doi": deposit["doi"], "citation_pdf_url": BASE + stem("en") + ".pdf",
                "citation_technical_report_number": REPORT, "citation_technical_report_institution": "Independent researcher", "citation_language": "en"}
    if lang == "en":
        meta = "\n".join(f'<meta name="{k}" content="{html.escape(v, quote=True)}">' for k, v in metadata.items())
    else:
        meta = '<meta name="robots" content="noindex,follow">\n<meta name="description" content="中文论证说明，不是全文翻译。Chinese companion guide, not a full translation or a separate paper.">'
    parts = ["<header>"]
    front, refs = True, False
    for i, (level, content) in enumerate(blocks):
        text = inline(content)
        if i == 0:
            parts.append("<h1>" + text + "</h1>")
        elif i == 1 and level == 2:
            parts.append('<p class="subtitle">' + text + "</p>")
        elif level:
            if front:
                parts.append("</header>")
                front = False
            refs = content == "References"
            tag = "h2" if content in ("Abstract", "摘要") or level == 2 else "h3"
            parts.append(f"<{tag}>{text}</{tag}>")
        else:
            cls = ' class="frontmatter"' if front else ' class="reference"' if refs else ""
            parts.append(f"<p{cls}>{text}</p>")
    if front:
        parts.append("</header>")
    nav = (f'<a href="{BASE}">English paper</a> · <a href="{BASE}zh-guide.html">中文论证说明</a> · '
           f'<a href="{BASE}{stem(lang)}.pdf">PDF</a> · <a href="{BASE}{stem(lang)}.docx">Word</a> · '
           f'<a href="https://doi.org/{deposit["doi"]}">DOI record</a> · <a href="{BASE}publication-record.json">Publication receipt</a>')
    return (f'<!doctype html>\n<html lang="{"en" if lang == "en" else "zh-CN"}"><head><meta charset="utf-8">'
            f'<meta name="viewport" content="width=device-width,initial-scale=1"><title>{html.escape(title)}</title>\n{meta}\n'
            f'<link rel="canonical" href="{canonical}"><style>{CSS}</style></head>\n<body><nav aria-label="Publication formats">{nav}</nav>'
            '<main><article>\n' + "\n".join(parts) + "\n</article></main></body></html>\n")


def citations(deposit, published):
    doi = deposit["doi"]
    url = "https://doi.org/" + doi
    csl = {"id": REPORT, "type": "report", "title": TITLE, "author": [{"family": "Liu", "given": "Hongju"}],
           "issued": {"date-parts": [[2026, 9, 19]]}, "DOI": doi, "URL": url,
           "publisher": "Zenodo", "number": REPORT, "version": VERSION, "language": "en",
           "note": "Conceptual preprint; not peer reviewed. Chinese companion is a guide, not a full translation."}
    write_json(published / "citation.csl.json", csl)
    (published / "citation.bib").write_text(
        "@techreport{Liu2026AIClaimant,\n  author = {Liu, Hongju},\n  title = {" + TITLE + "},\n  year = {2026},\n  institution = {Zenodo},\n  number = {" + REPORT + "},\n  version = {" + VERSION + "},\n  doi = {" + doi + "},\n  url = {" + url + "},\n  note = {Preprint; not peer reviewed}\n}\n", encoding="utf-8")
    (published / "citation.ris").write_text("TY  - RPRT\nTI  - " + TITLE + "\nAU  - Liu, Hongju\nPY  - 2026\nDA  - 2026/09/19\nPB  - Zenodo\nM1  - " + REPORT + "\nET  - " + VERSION + "\nDO  - " + doi + "\nUR  - " + url + "\nN1  - Preprint; not peer reviewed\nER  - \n", encoding="utf-8")


def build(deposit):
    published = ROOT / "published"
    published.mkdir(exist_ok=True)
    original = documents.create_base

    def styled(language):
        doc = original(language)
        doc.core_properties.subject = TITLE
        doc.core_properties.keywords = "scientific understanding, epistemic dependence, artificial consciousness"
        doc.styles["Reference"].paragraph_format.space_after = Pt(2)
        return doc

    documents.create_base = styled
    for lang, (source, _) in LANGS.items():
        md = published / (stem(lang) + ".md")
        md.write_text(publication_source(ROOT / "manuscript" / source, deposit), encoding="utf-8")
        blocks = documents.read_blocks(md)
        if lang == "en" and blocks[0][1] + ": " + blocks[1][1] != TITLE:
            raise ValueError("Manuscript title differs from reserved identity")
        documents.build(md, published / (stem(lang) + ".docx"), lang.upper())
        (published / (stem(lang) + ".html")).write_text(render_html(md, deposit, lang), encoding="utf-8")
    citations(deposit, published)
    print("Built source, Word, HTML and citations. Render both Word files before finalization.")


def check_formats(deposit):
    published, results = ROOT / "published", {}
    for lang in LANGS:
        name = stem(lang)
        blocks = documents.read_blocks(published / (name + ".md"))
        source = "\n".join(plain(text) for _, text in blocks)
        doc = Document(published / (name + ".docx"))
        word = "\n".join("".join(p._p.xpath(".//w:t/text()")) for p in doc.paragraphs)
        assert normalized(source) == normalized(word), (lang, "Source/Word text mismatch")
        parser = TextOnly(article_only=True)
        html_text = (published / (name + ".html")).read_text()
        parser.feed(html_text)
        assert normalized(source) == normalized("".join(parser.parts)), (lang, "Source/HTML mismatch")
        head = HeadMetadata()
        head.feed(html_text)
        assert head.canonicals == [BASE if lang == "en" else BASE + "zh-guide.html"]
        if lang == "en":
            for key, value in {"citation_title": TITLE, "citation_author": "Hongju Liu",
                               "citation_doi": deposit["doi"], "citation_publication_date": DATE.replace("-", "/"),
                               "citation_technical_report_number": REPORT, "citation_language": "en",
                               "citation_pdf_url": BASE + name + ".pdf"}.items():
                assert head.meta.get(key) == [value], ("HTML scholarly metadata", key)
        else:
            assert not any(key.startswith("citation_") for key in head.meta)
        with fitz.open(published / (name + ".pdf")) as pdf:
            text_parts = []
            for page in pdf:
                page_parts = []
                for block in page.get_text("blocks"):
                    x0, y0, x1, y1, text, _, kind = block
                    if kind != 0:
                        continue
                    assert x0 >= 0 and y0 >= 0 and x1 <= page.rect.width + 1 and y1 <= page.rect.height + 1
                    if y0 < page.rect.height - 40:
                        page_parts.append(text)
                assert len("".join(page_parts).strip()) > 80
                text_parts += page_parts
            pdf_text = "".join(text_parts)
            assert normalized(source) == normalized(pdf_text), (lang, "Source/PDF text mismatch")
            results[lang] = {"pages": len(pdf), "source_word_html_pdf_text_match": True, "page_boundary_check": True}
        if lang == "en":
            raw_refs = [text for _, text in blocks[blocks.index((2, "References")) + 1:]]
            source_refs = [plain(text) for text in raw_refs]
            word_refs = ["".join(p._p.xpath(".//w:t/text()")) for p in doc.paragraphs if p.style.name == "Reference"]
            assert len(source_refs) == len(word_refs) == 21
            assert [normalized(x) for x in source_refs] == [normalized(x) for x in word_refs]
            assert all("https://" in ref for ref in raw_refs)
            reference_urls = set()
            for ref in raw_refs:
                for match in documents.INLINE_RE.finditer(ref):
                    if match.group(2):
                        reference_urls.add(match.group(2))
                    elif match.group(3):
                        reference_urls.add(documents.split_url_punctuation(match.group(3))[0])
            word_urls = {rel.target_ref for rel in doc.part.rels.values() if rel.is_external}
            assert reference_urls <= word_urls, "A reference hyperlink was lost in Word"
            html_text = (published / (name + ".html")).read_text()
            assert all('href="' + html.escape(url, quote=True) + '"' in html_text for url in reference_urls)
    guide = (published / (stem("zh") + ".md")).read_text()
    guide_html = (published / (stem("zh") + ".html")).read_text()
    assert "不是全文翻译" in guide and 'content="noindex,follow"' in guide_html and "citation_title" not in guide_html
    csl = json.loads((published / "citation.csl.json").read_text())
    assert (csl["title"], csl["DOI"], csl["number"], csl["version"]) == (TITLE, deposit["doi"], REPORT, VERSION)
    for suffix in ("bib", "ris"):
        text = (published / ("citation." + suffix)).read_text()
        assert TITLE in text and deposit["doi"] in text and REPORT in text
    return {"source_docx_pdf_text_equal": True, "english_references_preserved_across_formats": True,
            "guide_explicitly_not_full_translation": True, "valid_citation_metadata": True,
            "pdf_text_extractable": True, "reference_entries": 21, "documents": results,
            "text_comparison": "Typography-normalized text: all letters and digits retained in order; whitespace and punctuation ignored",
            "reference_hyperlinks_preserved": True, "html_scholarly_metadata_checked": True,
            "peer_reviewed": False, "google_scholar_indexing": "NOT_ASSERTED"}


def finalize(deposit, visual_confirmed):
    if not visual_confirmed:
        raise ValueError("Actual completed page inspection must be confirmed")
    formats = check_formats(deposit)
    published = ROOT / "published"
    names = {p.name for p in published.iterdir()}
    if names - ALLOWED_FILES or ALLOWED_FILES - names - {"SHA256SUMS.txt"}:
        raise ValueError("Unexpected/missing publication files")
    sums = "\n".join(f"{digest(published / name)}  {name}" for name in sorted(ALLOWED_FILES - {"SHA256SUMS.txt"})) + "\n"
    (published / "SHA256SUMS.txt").write_text(sums)
    rows = [{"name": name, "bytes": (published / name).stat().st_size, "sha256": digest(published / name)} for name in sorted(ALLOWED_FILES)]
    expected = {key: deposit[key] for key in ("record_id", "doi", "title", "report_number", "version")}
    expected.update(file_count=len(rows), files=rows)
    write_json(ROOT / "EXPECTED-PUBLICATION.json", expected)
    write_json(ROOT / "format-checks.json", formats)
    write_json(ROOT / "visual-review.json", {"state": "VISUAL_AND_CONTENT_REVIEW_PASS", "expected_manifest_sha256": digest(ROOT / "EXPECTED-PUBLICATION.json"),
        "scope": "All pages of final English paper and Chinese guide inspected as rendered PNGs; contents and citations cross-checked.",
        "documents": formats["documents"], "review_method": "AI-assisted argument, source and page review; not independent scholarly peer review", "peer_reviewed": False})
    print(json.dumps(formats, indent=2))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("action", choices=("build", "check", "finalize"))
    parser.add_argument("--visual-review-confirmed", action="store_true")
    args = parser.parse_args()
    deposit = json.loads((ROOT / "deposit.json").read_text())
    validate_identity(deposit)
    if args.action == "build":
        build(deposit)
    elif args.action == "check":
        print(json.dumps(check_formats(deposit), indent=2))
    else:
        finalize(deposit, args.visual_review_confirmed)


if __name__ == "__main__":
    main()
