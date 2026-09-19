#!/usr/bin/env python3
"""Reusable typography and Markdown parsing helpers for the manuscript builder."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


LATIN_FONT = "Nimbus Roman"
CJK_FONT = "Noto Serif CJK SC"
BLACK = RGBColor(0, 0, 0)
URL_RE = re.compile(r"https?://[^\s<>]+")
INLINE_RE = re.compile(
    r"\[([^\]]+)\]\((https?://[^\s)]+)\)|"
    r"(https?://[^\s<>]+)|"
    r"\*\*([^*]+)\*\*|(?<!\*)\*([^*]+)\*(?!\*)|`([^`]+)`"
)


def font_properties(rpr, *, size=None, bold=None, italic=None, language="en-US"):
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for key in ("ascii", "hAnsi", "cs"):
        fonts.set(qn("w:" + key), LATIN_FONT)
    fonts.set(qn("w:eastAsia"), CJK_FONT)
    for key in ("asciiTheme", "hAnsiTheme", "cstheme", "eastAsiaTheme"):
        fonts.attrib.pop(qn("w:" + key), None)
    color = rpr.find(qn("w:color"))
    if color is None:
        color = OxmlElement("w:color")
        rpr.append(color)
    color.set(qn("w:val"), "000000")
    color.attrib.pop(qn("w:themeColor"), None)
    for prop, value in (("b", bold), ("i", italic)):
        if value is not None:
            element = rpr.find(qn("w:" + prop))
            if element is None:
                element = OxmlElement("w:" + prop)
                rpr.append(element)
            element.set(qn("w:val"), "1" if value else "0")
    if size is not None:
        for prop in ("sz", "szCs"):
            element = rpr.find(qn("w:" + prop))
            if element is None:
                element = OxmlElement("w:" + prop)
                rpr.append(element)
            element.set(qn("w:val"), str(round(size * 2)))
    lang = rpr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        rpr.append(lang)
    lang.set(qn("w:val"), language)
    lang.set(qn("w:eastAsia"), "zh-CN")


def no_border(element):
    for border in list(element.findall(qn("w:pBdr"))):
        element.remove(border)


def configure_style(style, *, size, before=0, after=6, spacing=1.12,
                    bold=False, italic=False, keep_next=False, language="en-US"):
    style.font.name = LATIN_FONT
    style.font.size = Pt(size)
    style.font.color.rgb = BLACK
    style.font.bold = bold
    style.font.italic = italic
    font_properties(style.element.get_or_add_rPr(), size=size, bold=bold,
                    italic=italic, language=language)
    pf = style.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = spacing
    pf.keep_with_next = keep_next
    pf.keep_together = False
    pf.widow_control = True
    no_border(style.element.get_or_add_pPr())


def create_base(language):
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.85)
    section.bottom_margin = Inches(0.85)
    section.left_margin = Inches(0.95)
    section.right_margin = Inches(0.95)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    zh = language == "ZH"
    lang = "zh-CN" if zh else "en-US"
    size = 11 if zh else 11.5
    spacing = 1.28 if zh else 1.13
    style_specs = {
        "Normal": dict(size=size, spacing=spacing, after=6),
        "Title": dict(size=19 if zh else 18, spacing=1.04, after=7,
                      bold=True, keep_next=True),
        "Subtitle": dict(size=12.5, spacing=1.10, after=12,
                         italic=not zh, keep_next=True),
        "Heading 1": dict(size=13 if zh else 13.5, before=13, after=6,
                          spacing=1.12, bold=True, keep_next=True),
        "Heading 2": dict(size=11.5, before=10, after=5,
                          spacing=1.12, bold=True, keep_next=True),
        "Heading 3": dict(size=11.5, before=8, after=4,
                          spacing=1.12, bold=True, keep_next=True),
        "Author": dict(size=11.5, after=5, spacing=1.10, keep_next=True),
        "Manuscript Metadata": dict(size=10, after=11, spacing=1.10,
                                    keep_next=True),
        "Keywords": dict(size=10.5, after=8, spacing=1.10),
        "Reference": dict(size=10.5, after=7, spacing=1.08),
    }
    for name, spec in style_specs.items():
        if name not in doc.styles:
            doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        # w:val describes the Latin-script language; w:eastAsia (set by
        # font_properties) describes Chinese. Using zh-CN for both makes
        # LibreOffice apply Chinese character-level breaks to Latin words.
        configure_style(doc.styles[name], language="en-US", **spec)
    doc.styles["Reference"].paragraph_format.left_indent = Inches(0.28)
    doc.styles["Reference"].paragraph_format.first_line_indent = Inches(-0.28)
    for name in ("Title", "Subtitle", "Author", "Manuscript Metadata"):
        doc.styles[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    # Do not inherit theme colors, border residue, or hanging title rules.
    for name in ("Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        no_border(doc.styles[name].element.get_or_add_pPr())

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    footer.paragraph_format.space_after = Pt(0)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    font_properties(rpr, size=9.5, language="en-US")
    run.append(rpr)
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    footer._p.append(field)

    settings = doc.settings.element
    update = OxmlElement("w:updateFields")
    update.set(qn("w:val"), "true")
    settings.append(update)
    hyphen = OxmlElement("w:autoHyphenation")
    hyphen.set(qn("w:val"), "false")
    settings.append(hyphen)
    doc.core_properties.author = "Hongju Liu"
    doc.core_properties.last_modified_by = "Hongju Liu"
    doc.core_properties.language = lang
    doc.core_properties.subject = "Scientific understanding and artificial consciousness"
    doc.core_properties.comments = ""
    return doc


def split_url_punctuation(url):
    # Punctuation outside the link is preserved as ordinary text. Balanced
    # parentheses within a URL remain in the link.
    suffix = ""
    while url and url[-1] in ".,;，。；":
        suffix = url[-1] + suffix
        url = url[:-1]
    while url.endswith(")") and url.count(")") > url.count("("):
        suffix = ")" + suffix
        url = url[:-1]
    return url, suffix


def add_hyperlink(paragraph, label, url, language):
    rel = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), rel)
    link.set(qn("w:history"), "1")
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    font_properties(props, language="en-US")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    proof = OxmlElement("w:noProof")
    props.append(proof)
    run.append(props)
    text = OxmlElement("w:t")
    # The target is always the complete original URL. Insert optional display
    # breaks only for a pathological unbroken component longer than 45 chars.
    # Current manuscript URLs have natural separators and need no intervention.
    display = re.sub(r"([^/:.?=&_\-]{45})(?=[^/:.?=&_\-])", lambda m: m.group(1) + "\u200b", label) \
        if any(len(s) > 45 for s in re.split(r"[/:.?=&_\-]", label)) else label
    text.text = display
    text.set(qn("xml:space"), "preserve")
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def add_inline(paragraph, content, language):
    cursor = 0
    for match in INLINE_RE.finditer(content):
        if match.start() > cursor:
            paragraph.add_run(content[cursor:match.start()])
        md_label, md_url, raw_url, bold, italic, code = match.groups()
        if md_url:
            add_hyperlink(paragraph, md_label, md_url, language)
        elif raw_url:
            url, suffix = split_url_punctuation(raw_url)
            add_hyperlink(paragraph, url, url, language)
            if suffix:
                paragraph.add_run(suffix)
        else:
            run = paragraph.add_run(bold or italic or code)
            if bold:
                run.bold = True
            if italic:
                run.italic = True
        cursor = match.end()
    if cursor < len(content):
        paragraph.add_run(content[cursor:])


def read_blocks(path):
    chunks = re.split(r"\n\s*\n", path.read_text(encoding="utf-8-sig").strip())
    blocks = []
    for chunk in chunks:
        lines = chunk.strip().splitlines()
        if any(line.startswith("```") for line in lines):
            raise ValueError(f"Unexpected fenced code in prose manuscript: {path}")
        if any(re.match(r"\s*\|.*\|\s*$", line) for line in lines):
            raise ValueError(f"Unexpected table in prose manuscript: {path}")
        # A heading is its own paragraph even if the author omitted a blank.
        prose = []
        for line in lines:
            heading = re.match(r"^(#{1,6})\s+(.+?)\s*#*\s*$", line)
            if heading:
                if prose:
                    blocks.append((0, " ".join(prose)))
                    prose = []
                blocks.append((len(heading.group(1)), heading.group(2)))
            else:
                # Convert simple list prefixes to normal prose rather than
                # leave Markdown markers visible. Current papers use prose.
                prose.append(re.sub(r"^\s*[-*+]\s+", "", line.strip()))
        if prose:
            blocks.append((0, " ".join(prose)))
    return blocks


def build(source, output, language):
    blocks = read_blocks(source)
    if not blocks or blocks[0][0] != 1:
        raise ValueError(f"Manuscript must start with one # title: {source}")
    doc = create_base(language)
    title = blocks[0][1]
    doc.core_properties.title = title
    frontmatter = True
    references = False
    frontmatter_prose = 0
    for index, (level, content) in enumerate(blocks):
        if level == 1:
            if index:
                raise ValueError(f"Only one level-one title is allowed: {source}")
            style = "Title"
        elif index == 1 and level == 2:
            style = "Subtitle"
        elif level:
            is_abstract = content.strip().casefold() in ("abstract", "摘要")
            if is_abstract:
                style = "Heading 1"
            else:
                style = f"Heading {min(max(level - 1, 1), 3)}"
            frontmatter = False
            references = content.strip().casefold() in ("references", "参考文献")
        elif frontmatter:
            frontmatter_prose += 1
            style = "Author" if frontmatter_prose == 1 else "Manuscript Metadata"
        elif content.startswith(("Keywords:", "Keywords：", "关键词：", "关键词:")):
            style = "Keywords"
        elif references:
            style = "Reference"
        else:
            style = "Normal"
        p = doc.add_paragraph(style=style)
        if language == "ZH":
            wrap = OxmlElement("w:wordWrap")
            wrap.set(qn("w:val"), "1")
            p._p.get_or_add_pPr().append(wrap)
        if level and content.strip().casefold() in ("references", "参考文献"):
            p.paragraph_format.page_break_before = True
        add_inline(p, content, language)
        no_border(p._p.get_or_add_pPr())
    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(f"Created {output} ({len(doc.paragraphs)} paragraphs)")

